from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

OUTPUT_DIR = Path(__file__).parent.parent / "output"


def _get_or_create(filename: str) -> Document:
    path = OUTPUT_DIR / filename
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if path.exists():
        return Document(str(path))
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def save_doc(doc: Document, filename: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    doc.save(str(path))
    return path


def write_captions_doc(captions: list[tuple[int, str]]) -> Path:
    """captions: list of (figure_number, caption_text)"""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    heading = doc.add_heading("Figure Captions", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for num, caption_text in captions:
        p = doc.add_paragraph()
        run_label = p.add_run(f"Figure {num}. ")
        run_label.bold = True
        p.add_run(caption_text.strip())
        p.paragraph_format.space_after = Pt(12)

    return save_doc(doc, "figures_captions.docx")


def append_section(filename: str, heading_text: str, body_text: str, level: int = 1) -> Path:
    doc = _get_or_create(filename)
    doc.add_heading(heading_text, level=level)
    # Split by double newline for paragraphs
    for para in body_text.strip().split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    doc.add_paragraph()  # spacing
    return save_doc(doc, filename)


def prepend_section(filename: str, heading_text: str, body_text: str) -> Path:
    """Insert a section at the beginning of the document (for Abstract/Title)."""
    existing = _get_or_create(filename)
    new_doc = Document()
    style = new_doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    new_doc.add_heading(heading_text, level=1)
    for para in body_text.strip().split("\n\n"):
        para = para.strip()
        if para:
            new_doc.add_paragraph(para)
    new_doc.add_paragraph()

    # Copy existing content
    for elem in existing.element.body:
        new_doc.element.body.append(elem)

    return save_doc(new_doc, filename)


def set_title(filename: str, title: str) -> Path:
    doc = _get_or_create(filename)
    # Insert title paragraph at beginning
    from docx.oxml.ns import qn
    from copy import deepcopy
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    # Move to front
    doc.element.body.insert(0, title_para._element)
    return save_doc(doc, filename)


def copy_doc(src: str, dst: str) -> Path:
    import shutil
    src_path = OUTPUT_DIR / src
    dst_path = OUTPUT_DIR / dst
    shutil.copy2(str(src_path), str(dst_path))
    return dst_path


def replace_references_section(filename: str, references: list[str]) -> Path:
    """
    Find the 'References' heading in the document and replace all content
    beneath it with the new formatted list.  If no References section exists,
    fall through to append_section.
    """
    doc = _get_or_create(filename)
    paragraphs = doc.paragraphs
    body = doc.element.body

    # Locate the References heading
    ref_heading_idx = None
    for i, para in enumerate(paragraphs):
        text = para.text.strip()
        style = para.style.name
        if (style.startswith("Heading") and "References" in text) or text == "References":
            ref_heading_idx = i
            break

    if ref_heading_idx is None:
        # No existing section — just append
        refs_body = "\n\n".join(f"{i+1}. {r}" for i, r in enumerate(references))
        return append_section(filename, "References", refs_body)

    # Find where the next section begins (next heading) or end of document
    next_section_idx = len(paragraphs)
    for i in range(ref_heading_idx + 1, len(paragraphs)):
        if paragraphs[i].style.name.startswith("Heading") and paragraphs[i].text.strip():
            next_section_idx = i
            break

    # Remove all paragraphs between the heading and the next section
    to_remove = [
        paragraphs[i]._element
        for i in range(ref_heading_idx + 1, next_section_idx)
    ]
    for elem in to_remove:
        body.remove(elem)

    # Insert reformatted reference entries immediately after the heading
    heading_elem = paragraphs[ref_heading_idx]._element
    heading_pos = list(body).index(heading_elem)
    for i, ref in enumerate(references, 1):
        p = doc.add_paragraph(f"{i}. {ref}")
        p_elem = p._element
        body.remove(p_elem)
        body.insert(heading_pos + i, p_elem)

    return save_doc(doc, filename)


def save_references_md(references: list[str]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "references.md"
    lines = ["# References\n"]
    for i, ref in enumerate(references, 1):
        lines.append(f"{i}. {ref}")
    path.write_text("\n".join(lines))
    return path


def save_tone_md(tone_text: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "tone.md"
    path.write_text(tone_text)
    return path


def save_intro_bullets(bullets: list[str]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "intro_research_bullets.md"
    lines = ["# Introduction Research Bullets\n"]
    for i, b in enumerate(bullets, 1):
        lines.append(f"{i}. {b}")
    path.write_text("\n".join(lines))
    return path
