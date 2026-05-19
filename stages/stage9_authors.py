"""
Stage 9 — Authorship formatting.

Display order: First Author, [Other Authors by priority], Corresponding Author
Affiliation symbols (in order of first appearance): †  ‡  §  ¶  ‖  #  ††  ‡‡  §§  ¶¶
Co-first mark: ✦  (footnote: "These authors contributed equally.")
Corresponding mark: *  (footnote: "Corresponding author: email")
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = Path(__file__).parent.parent / "output"

AFFILIATION_SYMBOLS = ["†", "‡", "§", "¶", "‖", "#", "††", "‡‡", "§§", "¶¶", "‖‖", "##"]
CO_FIRST_SYMBOL = "✦"
CORRESPONDING_SYMBOL = "*"


def parse_author_line(line: str) -> dict:
    """Parse 'Last, First | Affiliation 1 | Affiliation 2' into author dict."""
    parts = [p.strip() for p in line.split("|")]
    name = parts[0].strip() if parts else ""
    affiliations = [p for p in parts[1:] if p.strip()]
    return {"name": name, "affiliations": affiliations}


def parse_corresponding_line(line: str) -> dict:
    """
    Parse 'Last, First | Affiliation 1 | Affiliation 2 | email'.
    Detects the email as any pipe-delimited segment containing '@'.
    If no '@' is found, treats the last segment as the email field.
    """
    parts = [p.strip() for p in line.split("|")]
    name = parts[0].strip() if parts else ""
    rest = parts[1:]

    email = ""
    affiliations = []
    for seg in rest:
        if "@" in seg and not email:
            email = seg
        else:
            affiliations.append(seg)

    # Fallback: if still no email, pull from last segment
    if not email and affiliations:
        email = affiliations.pop()

    return {"name": name, "affiliations": affiliations, "email": email}


def assign_affiliation_symbols(ordered_authors: list[dict]) -> dict[str, str]:
    """
    Assign symbols to unique affiliations in order of first appearance
    across the full author list (first → others → corresponding).
    Returns {affiliation_string: symbol}.
    """
    aff_to_sym: dict[str, str] = {}
    sym_idx = 0
    for author in ordered_authors:
        for aff in author.get("affiliations", []):
            norm = aff.strip()
            if norm and norm not in aff_to_sym:
                if sym_idx < len(AFFILIATION_SYMBOLS):
                    aff_to_sym[norm] = AFFILIATION_SYMBOLS[sym_idx]
                    sym_idx += 1
    return aff_to_sym


def format_author_block(
    first_author: dict,
    corresponding_author: dict,
    other_authors: list[dict],
    co_first_names: list[str],
    aff_to_sym: dict[str, str],
) -> tuple[str, str]:
    """
    Build (author_line, affiliation_block).

    Author order: first_author → other_authors (priority order) → corresponding_author
    If first == corresponding, they appear only once (first position) with both marks.
    """
    co_first_set = {n.strip().lower() for n in co_first_names}

    same_person = (
        first_author["name"].strip().lower() == corresponding_author["name"].strip().lower()
    )
    ordered = [first_author] + other_authors
    if not same_person:
        ordered.append(corresponding_author)

    author_parts = []
    for author in ordered:
        name = author["name"]
        name_lower = name.strip().lower()

        # Affiliation superscripts (in the order the author's affiliations appear)
        aff_syms = "".join(
            aff_to_sym.get(a.strip(), "") for a in author.get("affiliations", [])
        )

        # Co-first mark: first author always gets it if co_first_names is non-empty;
        # others get it only if their name is explicitly listed.
        if co_first_names:
            if author is first_author or name_lower in co_first_set:
                co_mark = CO_FIRST_SYMBOL
            else:
                co_mark = ""
        else:
            co_mark = ""

        # Corresponding mark
        is_corr = (author is corresponding_author) or (same_person and author is first_author)
        corr_mark = CORRESPONDING_SYMBOL if is_corr else ""

        superscripts = co_mark + aff_syms + corr_mark
        author_parts.append(f"{name}{superscripts}")

    author_line = ", ".join(author_parts)

    # ── Affiliation block ──────────────────────────────────────────────────
    sorted_affs = sorted(
        aff_to_sym.items(),
        key=lambda x: (
            AFFILIATION_SYMBOLS.index(x[1]) if x[1] in AFFILIATION_SYMBOLS else 99
        ),
    )
    lines = [f"{sym} {aff}" for aff, sym in sorted_affs]

    # Corresponding author footnote
    corr_name = first_author["name"] if same_person else corresponding_author["name"]
    email = corresponding_author.get("email", "")
    lines.append(f"\n{CORRESPONDING_SYMBOL} Corresponding author: {corr_name}. E-mail: {email}")

    # Co-first footnote
    if co_first_names:
        lines.append(f"{CO_FIRST_SYMBOL} These authors contributed equally to this work.")

    affiliation_block = "\n".join(lines)
    return author_line, affiliation_block


def insert_authors_into_doc(filename: str, author_line: str, affiliation_block: str) -> Path:
    """Prepend the author line and affiliation block immediately after the title."""
    doc_path = OUTPUT / filename
    if not doc_path.exists():
        raise FileNotFoundError(f"{filename} not found in output/")

    doc = Document(str(doc_path))
    existing = list(doc.paragraphs)

    new_doc = Document()
    style = new_doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Preserve existing title (first paragraph)
    if existing:
        p = new_doc.add_paragraph(existing[0].text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if existing[0].runs:
            p.runs[0].bold = True

    # Author line (centered)
    auth_para = new_doc.add_paragraph(author_line)
    auth_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Affiliation block (left-aligned)
    new_doc.add_paragraph(affiliation_block)
    new_doc.add_paragraph()  # spacing

    # Copy remaining content
    for para in existing[1:]:
        new_doc.add_paragraph(para.text)

    new_doc.save(str(doc_path))
    return doc_path
